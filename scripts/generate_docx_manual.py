import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_shading(cell, color_hex):
    """Applies background color to a table cell."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=120, right=120):
    """Sets inner padding for a table cell (in dxa: 20 dxa = 1 pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4"):
    """Applies clean borders around and inside the table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def make_row_header(row):
    """Marks a table row to repeat on every new page."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:tblHeader'))

def set_row_cant_split(row):
    """Prevents table row from being split across page boundaries."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:cantSplit'))

def add_callout(doc, text, callout_type="NOTE"):
    """Adds a stylish callout box with a colored left border."""
    border_color = "0284C7" if callout_type == "NOTE" else ("16A34A" if callout_type == "TIP" else "DC2626")
    bg_color = "F0F9FF" if callout_type == "NOTE" else ("F0FDF4" if callout_type == "TIP" else "FEF2F2")
    
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(7.0)
    set_cell_shading(cell, bg_color)
    set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
    
    # Custom left border
    tblPr = tbl._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="none"/>\n'
        f'  <w:bottom w:val="none"/>\n'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'  <w:insideH w:val="none"/>\n'
        f'  <w:insideV w:val="none"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    r_tag = p.add_run(f"[{callout_type}] ")
    r_tag.bold = True
    r_tag.font.name = "Calibri"
    r_tag.font.size = Pt(10)
    if callout_type == "NOTE":
        r_tag.font.color.rgb = RGBColor(2, 132, 199)
    elif callout_type == "TIP":
        r_tag.font.color.rgb = RGBColor(22, 163, 74)
    else:
        r_tag.font.color.rgb = RGBColor(220, 38, 38)
        
    r_txt = p.add_run(text)
    r_txt.font.name = "Calibri"
    r_txt.font.size = Pt(9.5)
    r_txt.font.color.rgb = RGBColor(30, 41, 59)
    
    # Spacer
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def add_code_block(doc, code_text):
    """Adds a formatted code snippet box."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(7.0)
    set_cell_shading(cell, "F8FAFC")
    set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
    
    tblPr = tbl._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>\n'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>\n'
        f'  <w:left w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>\n'
        f'  <w:right w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>\n'
        f'  <w:insideH w:val="none"/>\n'
        f'  <w:insideV w:val="none"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    
    r = p.add_run(code_text.strip())
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(15, 23, 42)
    
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def render_markdown_table(doc, header_row, data_rows, col_widths=None):
    """Renders a styled table with alternating row shading and repeatable header."""
    num_cols = len(header_row)
    tbl = doc.add_table(rows=0, cols=num_cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, color="CBD5E1", sz="6")
    
    # Header
    hdr = tbl.add_row()
    make_row_header(hdr)
    set_row_cant_split(hdr)
    for i, title in enumerate(header_row):
        c = hdr.cells[i]
        if col_widths and i < len(col_widths):
            c.width = col_widths[i]
        set_cell_shading(c, "0F172A")  # Deep slate navy
        set_cell_margins(c, top=80, bottom=80, left=100, right=100)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(title.replace('<br>', '\n'))
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(255, 255, 255)
        
    # Data Rows
    for r_idx, row_data in enumerate(data_rows):
        row = tbl.add_row()
        set_row_cant_split(row)
        bg = "FFFFFF" if r_idx % 2 == 0 else "F8FAFC"
        for i, val in enumerate(row_data):
            c = row.cells[i]
            if col_widths and i < len(col_widths):
                c.width = col_widths[i]
            set_cell_shading(c, bg)
            set_cell_margins(c, top=60, bottom=60, left=90, right=90)
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.1
            
            clean_val = val.replace('<br>', '\n').replace('\\text{', '').replace('}', '').replace('$', '')
            clean_val = clean_val.replace('\\times', '×')
            
            # Simple formatting
            lines = clean_val.split('\n')
            for l_idx, line in enumerate(lines):
                if l_idx > 0:
                    p.add_run('\n')
                # Check bold tokens like **word**
                parts = re.split(r'(\*\*.*?\*\*)', line)
                for part in parts:
                    if part.startswith('**') and part.endswith('**'):
                        r = p.add_run(part[2:-2])
                        r.bold = True
                    else:
                        r = p.add_run(part)
                    r.font.name = "Calibri"
                    r.font.size = Pt(9)
                    r.font.color.rgb = RGBColor(30, 41, 59)
                    
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(2)
    sp.paragraph_format.space_after = Pt(6)

def generate_manual_docx():
    md_path = r'c:\Users\JV\Documents\GitHub\iReside\docs\4_USER_MANUAL_AND_INSTALLATION_GUIDE.md'
    out_repo = r'c:\Users\JV\Documents\GitHub\iReside\docs\iReside_User_Manual_and_Installation_Guide.docx'
    out_dl = r'C:\Users\JV\Downloads\iReside_User_Manual_and_Installation_Guide.docx'
    
    with open(md_path, 'r', encoding='utf-8') as f:
        md_text = f.read()
        
    doc = Document()
    
    # Page setup: Standard Letter Portrait with 0.75" margins
    for s in doc.sections:
        s.page_width = Inches(8.5)
        s.page_height = Inches(11.0)
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)
        
    # Header Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(2)
    r_ch = p_title.add_run("SECTION 4: USER MANUAL & INSTALLATION GUIDE\n")
    r_ch.bold = True
    r_ch.font.name = "Times New Roman"
    r_ch.font.size = Pt(14)
    r_ch.font.color.rgb = RGBColor(15, 23, 42)
    
    r_sub = p_title.add_run("iReside: Integrated Rental Property Management Platform\n")
    r_sub.bold = True
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(12)
    r_sub.font.color.rgb = RGBColor(30, 41, 59)
    
    r_meta = p_title.add_run("Academic Capstone Defense & Enterprise System Turnover Documentation | May 2026\n")
    r_meta.font.name = "Calibri"
    r_meta.font.size = Pt(9.5)
    r_meta.font.italic = True
    r_meta.font.color.rgb = RGBColor(100, 116, 139)
    
    # Divider line
    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_after = Pt(10)
    r_div = p_div.add_run("─" * 70)
    r_div.font.color.rgb = RGBColor(203, 213, 225)
    
    # Process lines
    lines = md_text.split('\n')
    idx = 0
    in_code = False
    code_buffer = []
    in_table = False
    table_header = []
    table_rows = []
    
    while idx < len(lines):
        line = lines[idx]
        stripped = line.strip()
        
        # Code block handling
        if stripped.startswith('```'):
            if in_code:
                in_code = False
                add_code_block(doc, '\n'.join(code_buffer))
                code_buffer = []
            else:
                in_code = True
                code_buffer = []
            idx += 1
            continue
            
        if in_code:
            code_buffer.append(line)
            idx += 1
            continue
            
        # Table handling
        if stripped.startswith('|') and stripped.endswith('|'):
            cells = [c.strip() for c in stripped.split('|')[1:-1]]
            # Check if delimiter row
            if all(re.match(r'^:?-+:?$', c) for c in cells):
                idx += 1
                continue
            if not in_table:
                in_table = True
                table_header = cells
                table_rows = []
            else:
                table_rows.append(cells)
            idx += 1
            continue
        else:
            if in_table:
                in_table = False
                render_markdown_table(doc, table_header, table_rows)
                table_header = []
                table_rows = []
                
        # Callout handling
        if stripped.startswith('> [!NOTE]'):
            callout_lines = []
            idx += 1
            while idx < len(lines) and lines[idx].strip().startswith('>'):
                callout_lines.append(lines[idx].strip()[1:].strip())
                idx += 1
            add_callout(doc, ' '.join(callout_lines), "NOTE")
            continue
        elif stripped.startswith('> [!TIP]'):
            callout_lines = []
            idx += 1
            while idx < len(lines) and lines[idx].strip().startswith('>'):
                callout_lines.append(lines[idx].strip()[1:].strip())
                idx += 1
            add_callout(doc, ' '.join(callout_lines), "TIP")
            continue
        elif stripped.startswith('> [!WARNING]') or stripped.startswith('> [!CAUTION]'):
            callout_lines = []
            idx += 1
            while idx < len(lines) and lines[idx].strip().startswith('>'):
                callout_lines.append(lines[idx].strip()[1:].strip())
                idx += 1
            add_callout(doc, ' '.join(callout_lines), "WARNING")
            continue
            
        # Headings
        if stripped.startswith('# PART'):
            doc.add_page_break()
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(16)
            h.paragraph_format.space_after = Pt(6)
            r = h.add_run(stripped.lstrip('#').strip())
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(14)
            r.font.color.rgb = RGBColor(15, 23, 42)
        elif stripped.startswith('## '):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            r = h.add_run(stripped[3:].strip())
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)
            r.font.color.rgb = RGBColor(30, 41, 59)
        elif stripped.startswith('### '):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            r = h.add_run(stripped[4:].strip())
            r.bold = True
            r.font.name = "Calibri"
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(51, 65, 85)
        elif stripped.startswith('#### '):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(6)
            h.paragraph_format.space_after = Pt(2)
            r = h.add_run(stripped[5:].strip())
            r.bold = True
            r.font.name = "Calibri"
            r.font.size = Pt(10.5)
            r.font.color.rgb = RGBColor(71, 85, 105)
        elif stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.15
            content = stripped[2:].strip()
            parts = re.split(r'(\*\*.*?\*\*)', content)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                else:
                    r = p.add_run(part)
                r.font.name = "Calibri"
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(30, 41, 59)
        elif re.match(r'^\d+\.\s', stripped):
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.15
            content = re.sub(r'^\d+\.\s', '', stripped)
            parts = re.split(r'(\*\*.*?\*\*)', content)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                else:
                    r = p.add_run(part)
                r.font.name = "Calibri"
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(30, 41, 59)
        elif stripped.startswith('---'):
            pass  # divider skipped or handled
        elif stripped == '':
            pass
        else:
            # Regular paragraph
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            parts = re.split(r'(\*\*.*?\*\*)', stripped)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                else:
                    r = p.add_run(part)
                r.font.name = "Calibri"
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(30, 41, 59)
                
        idx += 1
        
    if in_table:
        render_markdown_table(doc, table_header, table_rows)
        
    doc.save(out_repo)
    print(f"Successfully generated DOCX at: {out_repo}")
    
    try:
        doc.save(out_dl)
        print(f"Successfully copied DOCX to: {out_dl}")
    except Exception as e:
        print(f"Could not save copy to Downloads: {e}")

if __name__ == '__main__':
    generate_manual_docx()
