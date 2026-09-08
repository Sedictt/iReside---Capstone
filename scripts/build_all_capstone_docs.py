import os
import re
import shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def clean_latex_and_symbols(text):
    """Sanitizes raw LaTeX notations, formulas, and unwanted artifacts into clean Unicode."""
    if not text:
        return ""
    # m^3
    text = re.sub(r'\$\s*\\text\{m\}\^\{?3\}?\s*\$', 'm³', text)
    text = re.sub(r'\\text\{m\}\^\{?3\}?', 'm³', text)
    text = re.sub(r'\$\s*m\^3\s*\$', 'm³', text)
    text = re.sub(r'm\^3', 'm³', text)
    # arrows
    text = re.sub(r'\$\s*\\rightarrow\s*\$', ' → ', text)
    text = re.sub(r'\\rightarrow', ' → ', text)
    # math operators
    text = re.sub(r'\\times', '×', text)
    text = re.sub(r'\\pm', '±', text)
    text = re.sub(r'\\le', '≤', text)
    text = re.sub(r'\\ge', '≥', text)
    # fractions
    text = re.sub(r'\\frac\{([^}]+)\}\{([^}]+)\}', r'(\1 / \2)', text)
    # LaTeX text command
    text = re.sub(r'\\text\{([^}]+)\}', r'\1', text)
    # Strip double/single math dollar signs
    text = re.sub(r'\$\$([^\$]+)\$\$', r'\1', text)
    text = re.sub(r'\$([^\$]+)\$', r'\1', text)
    # Markdown links: [Title](url) -> Title
    text = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text)
    return text

def parse_inline_markdown(text):
    """Parses markdown bold, italic, and inline code tokens."""
    cleaned = clean_latex_and_symbols(text)
    pattern = r'(\*\*\*.*?\*\*\*|\*\*.*?\*\*|\*.*?\*|`.*?`)'
    tokens = re.split(pattern, cleaned)
    runs = []
    for t in tokens:
        if not t:
            continue
        if t.startswith('***') and t.endswith('***') and len(t) > 6:
            runs.append(('bold_italic', t[3:-3]))
        elif t.startswith('**') and t.endswith('**') and len(t) > 4:
            runs.append(('bold', t[2:-2]))
        elif t.startswith('*') and t.endswith('*') and len(t) > 2:
            runs.append(('italic', t[1:-1]))
        elif t.startswith('`') and t.endswith('`') and len(t) > 2:
            runs.append(('code', t[1:-1]))
        else:
            runs.append(('text', t))
    return runs

def add_formatted_runs(paragraph, text, base_font="Calibri", base_size=10, base_color=RGBColor(30, 41, 59)):
    """Applies parsed markdown inline tokens as native Word font properties."""
    tokens = parse_inline_markdown(text)
    for style_type, val in tokens:
        r = paragraph.add_run(val)
        if style_type == 'bold':
            r.bold = True
            r.font.name = base_font
            r.font.size = Pt(base_size)
            r.font.color.rgb = base_color
        elif style_type == 'italic':
            r.italic = True
            r.font.name = base_font
            r.font.size = Pt(base_size)
            r.font.color.rgb = base_color
        elif style_type == 'bold_italic':
            r.bold = True
            r.italic = True
            r.font.name = base_font
            r.font.size = Pt(base_size)
            r.font.color.rgb = base_color
        elif style_type == 'code':
            r.font.name = "Consolas"
            r.font.size = Pt(base_size - 0.5)
            r.font.color.rgb = RGBColor(15, 23, 42)
            r.bold = True
        else:
            r.font.name = base_font
            r.font.size = Pt(base_size)
            r.font.color.rgb = base_color

def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4"):
    tblPr = table._tbl.tblPr
    existing_borders = tblPr.find(qn('w:tblBorders'))
    if existing_borders is not None:
        tblPr.remove(existing_borders)
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def make_row_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:tblHeader'))

def set_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:cantSplit'))

def add_page_number_to_footer(footer, doc_subtitle="iReside Documentation"):
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    
    r_left = p.add_run(f"{doc_subtitle}  |  Page ")
    r_left.font.name = "Calibri"
    r_left.font.size = Pt(8.5)
    r_left.font.color.rgb = RGBColor(148, 163, 184)
    
    fld1 = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
    p._p.append(fld1)
    
    r_mid = p.add_run(" of ")
    r_mid.font.name = "Calibri"
    r_mid.font.size = Pt(8.5)
    r_mid.font.color.rgb = RGBColor(148, 163, 184)
    
    fld2 = parse_xml(r'<w:fldSimple %s w:instr="NUMPAGES"/>' % nsdecls('w'))
    p._p.append(fld2)

def add_header(header, doc_title="iReside Property Management Platform"):
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(doc_title)
    r.font.name = "Calibri"
    r.font.size = Pt(8.5)
    r.font.italic = True
    r.font.color.rgb = RGBColor(148, 163, 184)

def add_callout(doc, text, callout_type="NOTE"):
    border_color = "0284C7" if callout_type == "NOTE" else ("16A34A" if callout_type == "TIP" else "DC2626")
    bg_color = "F0F9FF" if callout_type == "NOTE" else ("F0FDF4" if callout_type == "TIP" else "FEF2F2")
    
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(7.0)
    set_cell_shading(cell, bg_color)
    set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
    
    tblPr = tbl._tbl.tblPr
    existing_borders = tblPr.find(qn('w:tblBorders'))
    if existing_borders is not None:
        tblPr.remove(existing_borders)
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="none"/>\n'
        f'  <w:bottom w:val="none"/>\n'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'  <w:insideH w:val="none"/>\n'
        f'  <w:insideV w:val="none"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    r_tag = p.add_run(f"[{callout_type}] ")
    r_tag.bold = True
    r_tag.font.name = "Calibri"
    r_tag.font.size = Pt(10)
    if callout_type == "NOTE":
        r_tag.font.color.rgb = RGBColor(2, 132, 199)
    elif callout_type == "TIP":
        r_tag.font.color.rgb = RGBColor(22, 163, 74)
    else:
        r_tag.font.color.rgb = RGBColor(220, 38, 38)
        
    add_formatted_runs(p, text, base_font="Calibri", base_size=9.5, base_color=RGBColor(30, 41, 59))
    
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(7.0)
    set_cell_shading(cell, "F8FAFC")
    set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
    
    tblPr = tbl._tbl.tblPr
    existing_borders = tblPr.find(qn('w:tblBorders'))
    if existing_borders is not None:
        tblPr.remove(existing_borders)
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>\n'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>\n'
        f'  <w:left w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>\n'
        f'  <w:right w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>\n'
        f'  <w:insideH w:val="none"/>\n'
        f'  <w:insideV w:val="none"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    
    r = p.add_run(code_text.strip())
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(15, 23, 42)
    
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def render_markdown_table(doc, header_row, data_rows, col_widths=None):
    num_cols = len(header_row)
    tbl = doc.add_table(rows=0, cols=num_cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, color="CBD5E1", sz="6")
    
    hdr = tbl.add_row()
    make_row_header(hdr)
    set_row_cant_split(hdr)
    for i, title in enumerate(header_row):
        c = hdr.cells[i]
        if col_widths and i < len(col_widths):
            c.width = col_widths[i]
        set_cell_shading(c, "0F172A")
        set_cell_margins(c, top=80, bottom=80, left=100, right=100)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        clean_title = clean_latex_and_symbols(title).replace('<br>', '\n')
        run = p.add_run(clean_title)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(255, 255, 255)
        
    for r_idx, row_data in enumerate(data_rows):
        row = tbl.add_row()
        set_row_cant_split(row)
        bg = "FFFFFF" if r_idx % 2 == 0 else "F8FAFC"
        for i, val in enumerate(row_data):
            c = row.cells[i]
            if col_widths and i < len(col_widths):
                c.width = col_widths[i]
            set_cell_shading(c, bg)
            set_cell_margins(c, top=60, bottom=60, left=90, right=90)
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.1
            
            clean_val = val.replace('<br>', '\n')
            lines = clean_val.split('\n')
            for l_idx, line in enumerate(lines):
                if l_idx > 0:
                    p.add_run('\n')
                add_formatted_runs(p, line, base_font="Calibri", base_size=9, base_color=RGBColor(30, 41, 59))
                    
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(2)
    sp.paragraph_format.space_after = Pt(6)

def compile_markdown_to_docx(md_path, out_path, title, subtitle):
    with open(md_path, 'r', encoding='utf-8') as f:
        md_text = f.read()
        
    doc = Document()
    
    for s in doc.sections:
        s.page_width = Inches(8.5)
        s.page_height = Inches(11.0)
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)
        add_header(s.header, title)
        add_page_number_to_footer(s.footer, subtitle)
        
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(14)
    p_title.paragraph_format.space_after = Pt(2)
    r_ch = p_title.add_run(f"{title.upper()}\n")
    r_ch.bold = True
    r_ch.font.name = "Times New Roman"
    r_ch.font.size = Pt(14)
    r_ch.font.color.rgb = RGBColor(15, 23, 42)
    
    r_sub = p_title.add_run("iReside: Integrated Rental Property Management Platform\n")
    r_sub.bold = True
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(12)
    r_sub.font.color.rgb = RGBColor(30, 41, 59)
    
    r_meta = p_title.add_run("Oral Defense Evaluation & System Turnover Edition | Academic Year 2025–2026\n")
    r_meta.font.name = "Calibri"
    r_meta.font.size = Pt(9.5)
    r_meta.font.italic = True
    r_meta.font.color.rgb = RGBColor(100, 116, 139)
    
    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_after = Pt(10)
    r_div = p_div.add_run("─" * 72)
    r_div.font.color.rgb = RGBColor(203, 213, 225)
    
    lines = md_text.split('\n')
    idx = 0
    in_code = False
    code_buffer = []
    in_table = False
    table_header = []
    table_rows = []
    
    while idx < len(lines):
        line = lines[idx]
        stripped = line.strip()
        
        # Code blocks
        if stripped.startswith('```'):
            if in_code:
                in_code = False
                add_code_block(doc, '\n'.join(code_buffer))
                code_buffer = []
            else:
                in_code = True
                code_buffer = []
            idx += 1
            continue
            
        if in_code:
            code_buffer.append(line)
            idx += 1
            continue
            
        # Tables
        if stripped.startswith('|') and stripped.endswith('|'):
            cells = [c.strip() for c in stripped.split('|')[1:-1]]
            if all(re.match(r'^:?-+:?$', c) for c in cells):
                idx += 1
                continue
            if not in_table:
                in_table = True
                table_header = cells
                table_rows = []
            else:
                table_rows.append(cells)
            idx += 1
            continue
        else:
            if in_table:
                in_table = False
                render_markdown_table(doc, table_header, table_rows)
                table_header = []
                table_rows = []
                
        # Callouts
        if stripped.startswith('> [!NOTE]'):
            callout_lines = []
            idx += 1
            while idx < len(lines) and lines[idx].strip().startswith('>'):
                callout_lines.append(lines[idx].strip()[1:].strip())
                idx += 1
            add_callout(doc, ' '.join(callout_lines), "NOTE")
            continue
        elif stripped.startswith('> [!TIP]'):
            callout_lines = []
            idx += 1
            while idx < len(lines) and lines[idx].strip().startswith('>'):
                callout_lines.append(lines[idx].strip()[1:].strip())
                idx += 1
            add_callout(doc, ' '.join(callout_lines), "TIP")
            continue
        elif stripped.startswith('> [!WARNING]') or stripped.startswith('> [!CAUTION]'):
            callout_lines = []
            idx += 1
            while idx < len(lines) and lines[idx].strip().startswith('>'):
                callout_lines.append(lines[idx].strip()[1:].strip())
                idx += 1
            add_callout(doc, ' '.join(callout_lines), "WARNING")
            continue
            
        # Headings
        if stripped.startswith('# PART') or stripped.startswith('# Part'):
            doc.add_page_break()
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(16)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
            cleaned_h = clean_latex_and_symbols(stripped.lstrip('#').strip())
            r = h.add_run(cleaned_h)
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(14)
            r.font.color.rgb = RGBColor(15, 23, 42)
        elif stripped.startswith('# '):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.keep_with_next = True
            cleaned_h = clean_latex_and_symbols(stripped[2:].strip())
            r = h.add_run(cleaned_h)
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(13)
            r.font.color.rgb = RGBColor(15, 23, 42)
        elif stripped.startswith('## '):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.keep_with_next = True
            cleaned_h = clean_latex_and_symbols(stripped[3:].strip())
            r = h.add_run(cleaned_h)
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)
            r.font.color.rgb = RGBColor(30, 41, 59)
        elif stripped.startswith('### '):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            h.paragraph_format.keep_with_next = True
            cleaned_h = clean_latex_and_symbols(stripped[4:].strip())
            r = h.add_run(cleaned_h)
            r.bold = True
            r.font.name = "Calibri"
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(51, 65, 85)
        elif stripped.startswith('#### '):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(6)
            h.paragraph_format.space_after = Pt(2)
            h.paragraph_format.keep_with_next = True
            cleaned_h = clean_latex_and_symbols(stripped[5:].strip())
            r = h.add_run(cleaned_h)
            r.bold = True
            r.font.name = "Calibri"
            r.font.size = Pt(10.5)
            r.font.color.rgb = RGBColor(71, 85, 105)
        elif stripped.startswith('- ') or stripped.startswith('* '):
            # Bullet list
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            content = stripped[2:].strip()
            add_formatted_runs(p, content, base_font="Calibri", base_size=10, base_color=RGBColor(30, 41, 59))
        elif re.match(r'^(\d+)\.\s+(.*)$', stripped):
            # Numbered list item: custom indent without global continuation bug
            m = re.match(r'^(\d+)\.\s+(.*)$', stripped)
            num = m.group(1)
            content = m.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            if content.endswith(':'):
                p.paragraph_format.keep_with_next = True
            r_num = p.add_run(f"{num}. ")
            r_num.bold = True
            r_num.font.name = "Calibri"
            r_num.font.size = Pt(10)
            r_num.font.color.rgb = RGBColor(30, 41, 59)
            add_formatted_runs(p, content, base_font="Calibri", base_size=10, base_color=RGBColor(30, 41, 59))
        elif stripped.startswith('---'):
            pass
        elif stripped == '':
            pass
        else:
            # Regular paragraph
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            if stripped.endswith(':'):
                p.paragraph_format.keep_with_next = True
            add_formatted_runs(p, stripped, base_font="Calibri", base_size=10, base_color=RGBColor(30, 41, 59))
                
        idx += 1
        
    if in_table:
        render_markdown_table(doc, table_header, table_rows)
        
    try:
        doc.save(out_path)
        print(f"Generated: {out_path}")
    except PermissionError:
        fallback_path = out_path.replace('.docx', '_Cleaned.docx')
        doc.save(fallback_path)
        print(f"File locked by Word! Saved updated version to: {fallback_path}")
        return fallback_path
    return out_path

def main():
    docs_dir = r'c:\Users\JV\Documents\GitHub\iReside\docs'
    dl_dir = r'C:\Users\JV\Downloads'
    
    sec4_md = os.path.join(docs_dir, '4_USER_MANUAL_AND_INSTALLATION_GUIDE.md')
    
    # Generate directly to a distinct fresh file so no permission error occurs
    final_docx = os.path.join(docs_dir, 'iReside_Section_4_User_Manual_and_Installation_Guide_FINAL.docx')
    compile_markdown_to_docx(sec4_md, final_docx, "Section 4: User Manual & Installation Guide", "iReside Master Manual")
    try:
        shutil.copy2(final_docx, os.path.join(dl_dir, 'iReside_Section_4_User_Manual_and_Installation_Guide_FINAL.docx'))
        print(f"Copied to Downloads: {os.path.join(dl_dir, 'iReside_Section_4_User_Manual_and_Installation_Guide_FINAL.docx')}")
    except Exception as e:
        print(f"Copy failed: {e}")
        
    sec4_docx = os.path.join(docs_dir, 'iReside_User_Manual_and_Installation_Guide.docx')
    saved_sec4 = compile_markdown_to_docx(sec4_md, sec4_docx, "Section 4: User Manual & Installation Guide", "iReside Master Manual")
    try:
        shutil.copy2(saved_sec4, os.path.join(dl_dir, 'iReside_User_Manual_and_Installation_Guide.docx'))
        print(f"Copied to Downloads: {os.path.join(dl_dir, 'iReside_User_Manual_and_Installation_Guide.docx')}")
    except Exception as e:
        print(f"Copy failed: {e}")
        
    um_md = os.path.join(docs_dir, 'USER_MANUAL.md')
    um_docx = os.path.join(docs_dir, 'iReside_User_Manual.docx')
    saved_um = compile_markdown_to_docx(um_md, um_docx, "iReside System User Manual", "User Manual")
    try:
        shutil.copy2(saved_um, os.path.join(dl_dir, 'iReside_User_Manual.docx'))
        print(f"Copied to Downloads: {os.path.join(dl_dir, 'iReside_User_Manual.docx')}")
    except Exception as e:
        print(f"Copy failed: {e}")
        
    ig_md = os.path.join(docs_dir, 'INSTALLATION_GUIDE.md')
    ig_docx = os.path.join(docs_dir, 'iReside_Installation_Guide.docx')
    saved_ig = compile_markdown_to_docx(ig_md, ig_docx, "iReside Installation & Deployment Guide", "Installation Guide")
    try:
        shutil.copy2(saved_ig, os.path.join(dl_dir, 'iReside_Installation_Guide.docx'))
        print(f"Copied to Downloads: {os.path.join(dl_dir, 'iReside_Installation_Guide.docx')}")
    except Exception as e:
        print(f"Copy failed: {e}")

if __name__ == '__main__':
    main()
